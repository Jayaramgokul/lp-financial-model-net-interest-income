"""04. Looking back: what the model's results owe to how it was built.

Five checks, each tied to a statement in the paper:
  1. Purchase fractions "do not affect NII": a property of two cells, not of the business.
  2. Collect 100% in the month of sale: a corner solution, because collecting faster costs
     nothing in the model. How large a cost would reverse it?
  3. The abstract calls the collection fractions decision variables. Making them LP
     variables gives the same answer in a single solve.
  4. "Multiple solutions for the same NII": is the optimum unique?
  5. The projected balance sheet: does it balance?

    python 04_looking_back.py
"""

import warnings

import numpy as np
import openpyxl
from scipy.optimize import linprog

from common import WORKBOOK
from model import SALES, TERM_RATES, Policy, Settings, affine, ending_cash, nii, receipts, solve

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
AS_RUN = Settings(rates=TERM_RATES.copy(), keep_d192_slip=True)


def lending_rows():
    A = np.zeros((6, 20))
    for m, cols in enumerate([(0, 6), (1, 7), (2, 8), (3, 9), (4,), (5,)]):
        A[m, list(cols)] = 1
    return A


def joint_lp(s, mcb=250):
    """One LP over the 20 instruments AND the collection fractions c0, c1 (c2 = 1 - c0 - c1)."""
    def unpack(v):
        return v[:20], Policy(cf=(v[20], v[21], 1 - v[20] - v[21]), mcb=mcb)
    c, c0 = affine(lambda v: np.atleast_1d(nii(*(lambda x, p: (x, p, s))(*unpack(v)))), n=22)
    Ae, be = affine(lambda v: ending_cash(unpack(v)[0], unpack(v)[1], s), n=22)
    Ar, br = affine(lambda v: receipts(unpack(v)[1]), n=22)           # receipts depend on c0, c1
    A_len = np.hstack([lending_rows(), np.zeros((6, 2))])
    A_ub = np.vstack([A_len - Ar, [[0] * 20 + [1, 1]]])                 # lending <= receipts; c0 + c1 <= 1
    b_ub = np.concatenate([br, [1]])
    r = linprog(-c.ravel(), A_ub=A_ub, b_ub=b_ub, A_eq=Ae, b_eq=mcb - be,
                bounds=[(0, None)] * 20 + [(0, 1), (0, 1)], method="highs")
    return float(c.ravel() @ r.x + c0[0]), r.x


def variable_ranges(p, s):
    """For each instrument, its smallest and largest value among all optimal plans."""
    z, _, _ = solve(p, s)
    c, c0 = affine(lambda x: np.atleast_1d(nii(x, p, s)))
    Ae, be = affine(lambda x: ending_cash(x, p, s))
    A_eq, b_eq = np.vstack([Ae, c]), np.concatenate([p.mcb - be, [z - c0[0]]])
    spans = []
    for i in range(20):
        lo_hi = []
        for sign in (1, -1):
            obj = np.zeros(20)
            obj[i] = sign
            r = linprog(obj, A_ub=lending_rows(), b_ub=receipts(p), A_eq=A_eq, b_eq=b_eq,
                        bounds=[(0, None)] * 20, method="highs")
            lo_hi.append(r.x[i])
        spans.append(lo_hi[1] - lo_hi[0])
    return z, np.array(spans)


def main():
    print("1. Payment fractions")
    for fix in (False, True):
        s = Settings(rates=TERM_RATES.copy(), purchase_timing_fix=fix)
        vals = [solve(Policy(pf=(t, 1 - t)), s)[0] for t in (0, 0.5, 1)]
        label = "purchases paid in the right month" if fix else "as in the workbook"
        print(f"   {label:<36} pay 0% / 50% / 100% in month of purchase: "
              + " / ".join(f"{v:.2f}" for v in vals))

    print("\n2. The value of collecting faster")
    base = solve(Policy(cf=(0.1, 0.1, 0.8)), AS_RUN)[0]
    best = solve(Policy(cf=(1.0, 0.0, 0.0)), AS_RUN)[0]
    sales = sum(SALES[:6])
    print(f"   company policy 10/10/80: {base:.2f}; everything in the month of sale: {best:.2f}")
    print(f"   gain {best - base:.2f} on six months' sales of {sales:,}: {100 * (best - base) / sales:.2f}% of sales")
    print("   so any cost of collecting at once (a cash discount, lost customers) above about 1% of sales")
    print("   would make the 'optimal' policy worse than the company's.")

    print("\n3. Collection fractions as LP decision variables (the abstract's formulation)")
    zj, v = joint_lp(AS_RUN)
    print(f"   one LP over 22 variables: NII {zj:.3f} with c0 = {abs(v[20]):.2f}, c1 = {abs(v[21]):.2f}, c2 = {abs(1 - v[20] - v[21]):.2f}")

    print("\n4. Is the optimum unique?")
    z, spans = variable_ranges(Policy(), AS_RUN)
    print(f"   NII {z:.4f}; largest range of any instrument across optimal plans: {spans.max():.2e}")

    print("\n5. The projected balance sheet (cached workbook values)")
    ws = openpyxl.load_workbook(WORKBOOK, data_only=True)["Sheet1"]
    for col, month in zip("BCDEFG", ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
        assets, claims = ws[f"{col}110"].value, ws[f"{col}119"].value
        print(f"   {month}: total assets {assets:9.2f}   liabilities + equity {claims:9.2f}   gap {claims - assets:9.2f}")

    print("\nWhat this shows:\n"
          "  1. The payment fractions cannot matter in the workbook, because both payment rows multiply the\n"
          "     same month's cost of goods. Paid in the right month, paying suppliers sooner costs NII, as it\n"
          "     should (-98.0 falls to -111.0).\n"
          "  2. Faster collection is worth about 1% of sales here, and the model charges nothing for it, so the\n"
          "     LP always picks the corner. A cost of collecting faster would turn this into a real trade-off.\n"
          "  3. Collection fractions enter the model linearly, so they can be optimised in the same LP. One\n"
          "     solve lands on the same corner that the paper's grid of Solver runs found.\n"
          "  4. The optimal plan is unique. The company's plan differs from it because it is not optimal\n"
          "     (-99.83 against -99.77 in the corrected model), not because there are several optima.\n"
          "  5. The balance sheet balances only in January. Its CD and note rows do not track outstanding\n"
          "     balances: the February gap of 680.89 is exactly the January 90-day CD, which drops out of\n"
          "     assets in February although it matures in April. The LP and NII do not use the balance\n"
          "     sheet, so the results are unaffected, but it could not serve as a check on the model.")


if __name__ == "__main__":
    main()
